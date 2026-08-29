"""Aggregate task-level benchmark runs and export a model-by-benchmark workbook."""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from ...benchmarks import BENCHMARK_STRUCTURE

METRIC_NAMES = ("Acc.", "Input.", "Output.", "Lat.")

# These tasks describe overlapping views of the same examples. The whole-benchmark
# task wins when it completed; otherwise the disjoint fallback tasks are combined.
OVERLAP_POLICIES: dict[str, tuple[tuple[str, ...], tuple[str, ...]]] = {
    "GAIA": (
        ("gaia_validation",),
        (
            "gaia_validation_level_1",
            "gaia_validation_level_2",
            "gaia_validation_level_3",
        ),
    ),
    "AFTraj-2K": (("aftraj_audit",), ("aftraj_audit_test",)),
}


@dataclass(frozen=True)
class TaskMetrics:
    model: str
    task: str
    cases: int
    score: float | None
    total_input_tokens: float | None
    total_output_tokens: float | None
    total_latency_s: float | None
    is_partial: bool
    path: Path

    @property
    def input_per_case(self) -> float | None:
        return _per_case(self.total_input_tokens, self.cases)

    @property
    def output_per_case(self) -> float | None:
        return _per_case(self.total_output_tokens, self.cases)

    @property
    def latency_per_case(self) -> float | None:
        return _per_case(self.total_latency_s, self.cases)


@dataclass(frozen=True)
class TaskRun:
    model: str
    task: str
    directory: Path
    expected_cases: int | None
    trial_count: int | None
    evaluation_count: int | None
    metrics: TaskMetrics | None

    @property
    def is_complete(self) -> bool:
        if self.metrics is None or self.metrics.is_partial:
            return False
        if self.expected_cases is not None and self.metrics.cases < self.expected_cases:
            return False
        return True


@dataclass(frozen=True)
class BenchmarkAggregate:
    model: str
    benchmark: str
    tasks: tuple[str, ...]
    cases: int
    score: float | None
    input_per_case: float | None
    output_per_case: float | None
    latency_per_case: float | None
    status: str
    note: str


def discover_task_runs(
    runs_root: str | Path,
    *,
    duplicate_policy: str = "error",
) -> list[TaskRun]:
    """Read task run directories below ``runs_root``.

    A task directory is recognized by ``config.yaml`` or ``metrics.json``. Metrics
    are preferred for model/task identity; resolved config values provide expected
    case counts and make interrupted runs visible even when metrics were never written.
    """

    root = Path(runs_root).expanduser().resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"benchmark runs root does not exist: {root}")

    directories = {path.parent for path in root.rglob("metrics.json")}
    directories.update(path.parent for path in root.rglob("config.yaml"))
    if not directories:
        raise FileNotFoundError(f"no metrics.json or config.yaml found below: {root}")

    runs: list[TaskRun] = []
    for directory in sorted(directories):
        metrics_data = _read_json(directory / "metrics.json")
        config_data = _read_yaml(directory / "config.yaml")
        resolved = config_data.get("resolved", {}) if isinstance(config_data, dict) else {}

        model = str(
            metrics_data.get("model") or resolved.get("model") or _model_from_path(directory, root)
        )
        task = str(metrics_data.get("task") or resolved.get("task") or directory.name)
        expected_cases = _positive_int(resolved.get("n"))
        from lychee_mas.eval.evaluation.projections import build_result_projection

        trial_rows = build_result_projection(directory)
        trials = len(trial_rows)
        evaluations = sum(row["evaluation_status"] == "completed" for row in trial_rows)
        metrics = (
            _parse_task_metrics(metrics_data, directory / "metrics.json", model, task)
            if metrics_data
            else None
        )
        runs.append(
            TaskRun(
                model=model,
                task=task,
                directory=directory,
                expected_cases=expected_cases,
                trial_count=trials,
                evaluation_count=evaluations,
                metrics=metrics,
            )
        )

    return _resolve_duplicates(runs, duplicate_policy)


def aggregate_benchmarks(
    runs: Sequence[TaskRun],
    *,
    benchmark_structure: Sequence[Mapping[str, Any]] = BENCHMARK_STRUCTURE,
) -> tuple[list[str], list[str], dict[tuple[str, str], BenchmarkAggregate]]:
    """Collapse runnable tasks into one weighted result per model and benchmark."""

    models = sorted({run.model for run in runs}, key=str.casefold)
    benchmarks = [str(row["benchmark_source"]) for row in benchmark_structure]
    by_model_task = {(run.model, run.task): run for run in runs}
    aggregates: dict[tuple[str, str], BenchmarkAggregate] = {}

    for model in models:
        for row in benchmark_structure:
            benchmark = str(row["benchmark_source"])
            expected_tasks = tuple(str(task) for task in row["runnable_tasks"])
            selected, complete, note = _select_runs(
                model,
                benchmark,
                expected_tasks,
                by_model_task,
            )
            metrics = [run.metrics for run in selected if run.metrics is not None]
            if not metrics:
                continue
            cases = sum(metric.cases for metric in metrics)
            aggregates[(model, benchmark)] = BenchmarkAggregate(
                model=model,
                benchmark=benchmark,
                tasks=tuple(metric.task for metric in metrics),
                cases=cases,
                score=_weighted_score(metrics, cases),
                input_per_case=_weighted_total(metrics, "total_input_tokens", cases),
                output_per_case=_weighted_total(metrics, "total_output_tokens", cases),
                latency_per_case=_weighted_total(metrics, "total_latency_s", cases),
                status="complete" if complete else "partial",
                note=note,
            )
    return models, benchmarks, aggregates


def build_task_detail_rows(
    runs: Sequence[TaskRun],
    *,
    benchmark_structure: Sequence[Mapping[str, Any]] = BENCHMARK_STRUCTURE,
) -> list[dict[str, Any]]:
    """Return audit-friendly rows for every expected runnable task."""

    models = sorted({run.model for run in runs}, key=str.casefold)
    by_model_task = {(run.model, run.task): run for run in runs}
    rows: list[dict[str, Any]] = []

    for model in models:
        for structure_row in benchmark_structure:
            benchmark = str(structure_row["benchmark_source"])
            expected = tuple(str(task) for task in structure_row["runnable_tasks"])
            selected, _, _ = _select_runs(model, benchmark, expected, by_model_task)
            selected_tasks = {run.task for run in selected}
            for task in expected:
                run = by_model_task.get((model, task))
                role, status, note = _task_detail_status(benchmark, task, run, selected_tasks)
                metric = run.metrics if run else None
                rows.append(
                    {
                        "model": model,
                        "benchmark": benchmark,
                        "task": task,
                        "aggregation_role": role,
                        "status": status,
                        "cases": metric.cases if metric else None,
                        "expected_cases": run.expected_cases if run else None,
                        "trials": run.trial_count if run else None,
                        "evaluations": run.evaluation_count if run else None,
                        "score": metric.score if metric else None,
                        "input_per_case": metric.input_per_case if metric else None,
                        "output_per_case": metric.output_per_case if metric else None,
                        "latency_per_case": metric.latency_per_case if metric else None,
                        "metrics_path": str(metric.path if metric else run.directory)
                        if run
                        else "",
                        "note": note,
                    }
                )
    return rows


def write_benchmark_workbook(
    runs_root: str | Path,
    output_path: str | Path,
    *,
    duplicate_policy: str = "error",
) -> Path:
    """Create the requested model-by-benchmark Excel workbook."""

    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - dependency error path
        raise RuntimeError(
            "Excel export requires openpyxl; install the LycheeMAS benchmark extra"
        ) from exc

    runs = discover_task_runs(runs_root, duplicate_policy=duplicate_policy)
    models, benchmarks, aggregates = aggregate_benchmarks(runs)
    detail_rows = build_task_detail_rows(runs)

    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    summary = wb.active
    summary.title = "Benchmark Summary"
    details = wb.create_sheet("Task Details")
    audit = wb.create_sheet("Audit")
    definitions = wb.create_sheet("Definitions")

    navy = "19324D"
    blue = "2F75B5"
    pale_blue = "DCEAF7"
    pale_green = "E2F0D9"
    pale_amber = "FFF2CC"
    pale_red = "FCE4D6"
    pale_gray = "E7E6E6"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="D9E1F2")

    summary.sheet_view.showGridLines = False
    summary.freeze_panes = "B3"
    summary.merge_cells("A1:A2")
    summary["A1"] = "Model"
    summary["A1"].fill = PatternFill("solid", fgColor=navy)
    summary["A1"].font = Font(color=white, bold=True)
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary.column_dimensions["A"].width = 31
    summary.row_dimensions[1].height = 28
    summary.row_dimensions[2].height = 24

    for benchmartrial_index, benchmark in enumerate(benchmarks):
        start_col = 2 + benchmartrial_index * len(METRIC_NAMES)
        end_col = start_col + len(METRIC_NAMES) - 1
        summary.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        top = summary.cell(1, start_col, benchmark)
        top.fill = PatternFill("solid", fgColor=blue)
        top.font = Font(color=white, bold=True)
        top.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(start_col, end_col + 1):
            cell = summary.cell(2, col, METRIC_NAMES[col - start_col])
            cell.fill = PatternFill("solid", fgColor=pale_blue)
            cell.font = Font(color=navy, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin_gray)
            summary.column_dimensions[get_column_letter(col)].width = 12

    for row_index, model in enumerate(models, start=3):
        model_cell = summary.cell(row_index, 1, model)
        model_cell.font = Font(bold=True, color=navy)
        model_cell.alignment = Alignment(vertical="center")
        model_cell.fill = PatternFill("solid", fgColor="F3F6F9")
        summary.row_dimensions[row_index].height = 24
        for benchmartrial_index, benchmark in enumerate(benchmarks):
            aggregate = aggregates.get((model, benchmark))
            start_col = 2 + benchmartrial_index * len(METRIC_NAMES)
            if aggregate is None:
                for col in range(start_col, start_col + len(METRIC_NAMES)):
                    summary.cell(row_index, col).fill = PatternFill("solid", fgColor=pale_gray)
                continue
            values = (
                aggregate.score,
                aggregate.input_per_case,
                aggregate.output_per_case,
                aggregate.latency_per_case,
            )
            for offset, value in enumerate(values):
                cell = summary.cell(row_index, start_col + offset, value)
                cell.alignment = Alignment(horizontal="right")
                cell.fill = PatternFill(
                    "solid", fgColor=pale_green if aggregate.status == "complete" else pale_amber
                )
            summary.cell(row_index, start_col).number_format = "0.00%"
            summary.cell(row_index, start_col + 1).number_format = "#,##0.0"
            summary.cell(row_index, start_col + 2).number_format = "#,##0.0"
            summary.cell(row_index, start_col + 3).number_format = "#,##0.000"
            summary.cell(row_index, start_col).comment = Comment(
                f"status={aggregate.status}; cases={aggregate.cases}; "
                f"tasks={', '.join(aggregate.tasks)}. {aggregate.note}",
                "LycheeMAS",
            )

    last_row = max(3, 2 + len(models))
    last_col = 1 + len(benchmarks) * len(METRIC_NAMES)
    summary.auto_filter.ref = f"A2:{get_column_letter(last_col)}{last_row}"
    if models:
        for benchmartrial_index in range(len(benchmarks)):
            col = 2 + benchmartrial_index * len(METRIC_NAMES)
            score_range = (
                summary.cell(3, col).coordinate + ":" + summary.cell(last_row, col).coordinate
            )
            summary.conditional_formatting.add(
                score_range,
                ColorScaleRule(
                    start_type="num",
                    start_value=0,
                    start_color="F8696B",
                    mid_type="num",
                    mid_value=0.5,
                    mid_color="FFEB84",
                    end_type="num",
                    end_value=1,
                    end_color="63BE7B",
                ),
            )

    detail_headers = [
        "Model",
        "Benchmark",
        "Runnable task",
        "Aggregation role",
        "Status",
        "Cases",
        "Expected cases",
        "Trials",
        "Evaluations",
        "Native score",
        "Input / case",
        "Output / case",
        "Latency / case (s)",
        "Metrics path",
        "Note",
    ]
    detail_keys = [
        "model",
        "benchmark",
        "task",
        "aggregation_role",
        "status",
        "cases",
        "expected_cases",
        "trials",
        "evaluations",
        "score",
        "input_per_case",
        "output_per_case",
        "latency_per_case",
        "metrics_path",
        "note",
    ]
    _write_tabular_sheet(details, detail_headers, detail_rows, detail_keys, navy, white)
    for row_index in range(2, 2 + len(detail_rows)):
        details.row_dimensions[row_index].height = 42
        details.cell(row_index, 10).number_format = "0.00%"
        details.cell(row_index, 11).number_format = "#,##0.0"
        details.cell(row_index, 12).number_format = "#,##0.0"
        details.cell(row_index, 13).number_format = "#,##0.000"
        status = details.cell(row_index, 5).value
        color = pale_green if status == "complete" else pale_amber
        if status in {"missing", "failed"}:
            color = pale_red
        details.cell(row_index, 5).fill = PatternFill("solid", fgColor=color)
    _set_widths(
        details,
        [28, 22, 30, 18, 14, 11, 15, 13, 11, 14, 15, 15, 18, 65, 60],
    )

    audit_rows = _audit_rows(detail_rows)
    audit_headers = ["Severity", "Model", "Benchmark", "Runnable task", "Message", "Path"]
    audit_keys = ["severity", "model", "benchmark", "task", "message", "path"]
    _write_tabular_sheet(audit, audit_headers, audit_rows, audit_keys, navy, white)
    _set_widths(audit, [12, 28, 22, 30, 90, 70])
    for row_index in range(2, 2 + len(audit_rows)):
        severity = audit.cell(row_index, 1).value
        color = pale_red if severity == "ERROR" else pale_amber
        audit.cell(row_index, 1).fill = PatternFill("solid", fgColor=color)

    definition_rows = [
        (
            "Acc.",
            "按 case 数加权的 benchmark 原生分数均值。二值/选择题为正确率；"
            "LoCoMo10、MAST 等保留其原生 F1 或分类分数。",
        ),
        (
            "Input.",
            "每个 case 的模型输入文本 token 数：纳入任务 "
            "total_input_text_tokens 之和 / case 数之和。",
        ),
        (
            "Output.",
            "每个 case 的模型输出文本 token 数：纳入任务 "
            "total_output_text_tokens 之和 / case 数之和。",
        ),
        (
            "Lat.",
            "每个 case 的模型生成延迟（秒）：纳入任务 total_model_latency_s "
            "之和 / case 数之和；不是端到端墙钟时间。",
        ),
        (
            "加权规则",
            "多个互不重叠 task 按 case 数合并。AgentCollabBench 合并四个子任务；"
            "GAIA 总集与 level 拆分二选一；AFTraj 全量与 test 子集二选一。",
        ),
        (
            "黄色单元格",
            "该模型在该 benchmark 上只有部分可评分结果；数值仅基于已有 "
            "metrics.json，具体缺失见 Audit。",
        ),
        ("来源目录", str(Path(runs_root).expanduser().resolve())),
    ]
    definitions.sheet_view.showGridLines = False
    definitions["A1"] = "Field"
    definitions["B1"] = "Definition"
    for cell in definitions[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True)
    for row_index, (field, definition) in enumerate(definition_rows, start=2):
        definitions.cell(row_index, 1, field).font = Font(bold=True, color=navy)
        definitions.cell(row_index, 2, definition).alignment = Alignment(wrap_text=True)
        definitions.row_dimensions[row_index].height = 34
    definitions.column_dimensions["A"].width = 22
    definitions.column_dimensions["B"].width = 110
    definitions.freeze_panes = "A2"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(output)
    return output


def _select_runs(
    model: str,
    benchmark: str,
    expected_tasks: tuple[str, ...],
    by_model_task: Mapping[tuple[str, str], TaskRun],
) -> tuple[list[TaskRun], bool, str]:
    if benchmark in OVERLAP_POLICIES:
        preferred, fallback = OVERLAP_POLICIES[benchmark]
        preferred_runs = [by_model_task.get((model, task)) for task in preferred]
        if preferred_runs and all(run is not None and run.is_complete for run in preferred_runs):
            selected = [run for run in preferred_runs if run is not None]
            return (
                selected,
                True,
                "Used the complete whole-benchmark task; overlapping subsets excluded.",
            )

        fallback_runs = [by_model_task.get((model, task)) for task in fallback]
        selected = [run for run in fallback_runs if run is not None and run.metrics is not None]
        complete = bool(fallback_runs) and all(
            run is not None and run.is_complete for run in fallback_runs
        )
        missing = [
            task
            for task, run in zip(fallback, fallback_runs)
            if run is None or run.metrics is None or not run.is_complete
        ]
        note = "Used disjoint fallback tasks because the whole-benchmark task was incomplete."
        if missing:
            note += f" Missing/incomplete: {', '.join(missing)}."
        return selected, complete, note

    task_runs = [by_model_task.get((model, task)) for task in expected_tasks]
    selected = [run for run in task_runs if run is not None and run.metrics is not None]
    complete = bool(task_runs) and all(run is not None and run.is_complete for run in task_runs)
    missing = [
        task
        for task, run in zip(expected_tasks, task_runs)
        if run is None or run.metrics is None or not run.is_complete
    ]
    note = "All expected runnable tasks were case-weighted."
    if missing:
        note += f" Missing/incomplete: {', '.join(missing)}."
    return selected, complete, note


def _task_detail_status(
    benchmark: str,
    task: str,
    run: TaskRun | None,
    selected_tasks: set[str],
) -> tuple[str, str, str]:
    if run is None:
        return "missing", "missing", "No run directory was found."
    if run.metrics is None:
        return (
            "missing",
            "missing",
            "The Run has no metrics.json; completed Trials have not been scored.",
        )
    if task not in selected_tasks:
        return (
            "excluded-overlap",
            "complete" if run.is_complete else "partial",
            "Excluded to prevent double-counting an overlapping whole set/subset.",
        )
    if run.is_complete:
        return "included", "complete", "Included in the benchmark-level weighted result."
    return "included", "partial", "Included, but the task run is incomplete."


def _parse_task_metrics(data: Mapping[str, Any], path: Path, model: str, task: str) -> TaskMetrics:
    cases = _positive_int(data.get("case_count") or data.get("num_cases"))
    if cases is None:
        raise ValueError(f"metrics file has no positive case count: {path}")
    score = _number(data.get("accuracy"))
    if score is None:
        score = _number(data.get("score_mean"))
    return TaskMetrics(
        model=model,
        task=task,
        cases=cases,
        score=score,
        total_input_tokens=_metric_total(
            data, "total_input_text_tokens", "mean_input_text_tokens_per_case", cases
        ),
        total_output_tokens=_metric_total(
            data, "total_output_text_tokens", "mean_output_text_tokens_per_case", cases
        ),
        total_latency_s=_metric_total(
            data, "total_model_latency_s", "mean_model_latency_s_per_case", cases
        ),
        is_partial=bool(data.get("is_partial", False)),
        path=path.resolve(),
    )


def _metric_total(
    data: Mapping[str, Any], total_name: str, mean_name: str, cases: int
) -> float | None:
    total = _number(data.get(total_name))
    if total is not None:
        return total
    mean = _number(data.get(mean_name))
    return mean * cases if mean is not None else None


def _weighted_score(metrics: Sequence[TaskMetrics], cases: int) -> float | None:
    scored = [(metric.score, metric.cases) for metric in metrics if metric.score is not None]
    scored_cases = sum(case_count for _, case_count in scored)
    if not scored or scored_cases <= 0:
        return None
    return sum(float(score) * case_count for score, case_count in scored) / scored_cases


def _weighted_total(metrics: Sequence[TaskMetrics], field: str, cases: int) -> float | None:
    if cases <= 0:
        return None
    values = [getattr(metric, field) for metric in metrics]
    if any(value is None for value in values):
        return None
    return sum(float(value) for value in values if value is not None) / cases


def _audit_rows(detail_rows: Iterable[Mapping[str, Any]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for row in detail_rows:
        role = str(row["aggregation_role"])
        status = str(row["status"])
        if role == "excluded-overlap" and status == "complete":
            continue
        if role == "included" and status == "complete":
            continue
        severity = "ERROR" if status == "missing" else "WARN"
        rows.append(
            {
                "severity": severity,
                "model": str(row["model"]),
                "benchmark": str(row["benchmark"]),
                "task": str(row["task"]),
                "message": str(row["note"]),
                "path": str(row["metrics_path"]),
            }
        )
    if not rows:
        rows.append(
            {
                "severity": "OK",
                "model": "",
                "benchmark": "",
                "task": "",
                "message": "No missing or partial benchmark task runs were detected.",
                "path": "",
            }
        )
    return rows


def _write_tabular_sheet(
    sheet: Any,
    headers: Sequence[str],
    rows: Sequence[Mapping[str, Any]],
    keys: Sequence[str],
    header_color: str,
    header_font_color: str,
) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(1, column, header)
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.font = Font(color=header_font_color, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, row in enumerate(rows, start=2):
        for column, key in enumerate(keys, start=1):
            cell = sheet.cell(row_index, column, row.get(key))
            cell.alignment = Alignment(vertical="top", wrap_text=column >= len(keys) - 1)
    last_row = max(1, len(rows) + 1)
    sheet.auto_filter.ref = f"A1:{sheet.cell(last_row, len(headers)).coordinate}"


def _set_widths(sheet: Any, widths: Sequence[float]) -> None:
    from openpyxl.utils import get_column_letter

    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def _resolve_duplicates(runs: Sequence[TaskRun], policy: str) -> list[TaskRun]:
    if policy not in {"error", "latest"}:
        raise ValueError("duplicate_policy must be 'error' or 'latest'")
    grouped: dict[tuple[str, str], list[TaskRun]] = {}
    for run in runs:
        grouped.setdefault((run.model, run.task), []).append(run)
    resolved: list[TaskRun] = []
    for key, candidates in grouped.items():
        if len(candidates) == 1:
            resolved.append(candidates[0])
            continue
        if policy == "error":
            paths = "\n  ".join(str(run.directory) for run in candidates)
            raise ValueError(f"duplicate runs for model/task {key}:\n  {paths}")
        resolved.append(max(candidates, key=lambda run: run.directory.stat().st_mtime))
    return sorted(resolved, key=lambda run: (run.model.casefold(), run.task, str(run.directory)))


def _read_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError(f"expected JSON object: {path}")
    return data


def _read_yaml(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    try:
        import yaml
    except ImportError as exc:  # pragma: no cover - dependency error path
        raise RuntimeError("benchmark run discovery requires PyYAML") from exc
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    return data if isinstance(data, dict) else {}


def _model_from_path(directory: Path, root: Path) -> str:
    relative = directory.relative_to(root)
    parts = relative.parts
    return parts[-3] if len(parts) >= 3 else "unknown-model"


def _positive_int(value: Any) -> int | None:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return None
    return parsed if parsed > 0 else None


def _number(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _per_case(total: float | None, cases: int) -> float | None:
    return total / cases if total is not None and cases > 0 else None
