import json

import yaml
from lychee_mas.eval.evaluation.reports.benchmark_excel import (
    aggregate_benchmarks,
    discover_task_runs,
    write_benchmark_workbook,
)
from lychee_mas.runtime.events.store import RunEventWriter, run_events_path


def _write_run(
    root,
    *,
    model,
    task,
    cases,
    score,
    input_tokens,
    output_tokens,
    latency_s,
    expected=None,
    with_metrics=True,
):
    directory = root / model / "none" / task
    directory.mkdir(parents=True)
    (directory / "config.yaml").write_text(
        yaml.safe_dump({"resolved": {"model": model, "task": task, "n": expected or cases}}),
        encoding="utf-8",
    )
    if with_metrics:
        (directory / "metrics.json").write_text(
            json.dumps(
                {
                    "model": model,
                    "task": task,
                    "case_count": cases,
                    "accuracy": score,
                    "total_input_text_tokens": input_tokens,
                    "total_output_text_tokens": output_tokens,
                    "total_model_latency_s": latency_s,
                    "is_partial": False,
                }
            ),
            encoding="utf-8",
        )
        writer = RunEventWriter(run_events_path(directory))
        for index in range(cases):
            record = {"case_id": str(index), "dataset_index": index, "trial_index": 0}
            started_id = writer.log_event("trial.started", **record)
            terminal_id = writer.record_trial(
                {**record, "operation_id": started_id, "final_output": "answer"}
            )
            writer.record_evaluation({**record, "trial_event_id": terminal_id, "score": score})
    return directory


def test_weighted_aggregation_and_overlap_deduplication(tmp_path):
    model = "Qwen3-4B-Instruct-2507"
    _write_run(
        tmp_path,
        model=model,
        task="agent_collab_idr",
        cases=2,
        score=0.5,
        input_tokens=20,
        output_tokens=10,
        latency_s=4,
    )
    _write_run(
        tmp_path,
        model=model,
        task="agent_collab_rtd",
        cases=6,
        score=1.0,
        input_tokens=180,
        output_tokens=90,
        latency_s=24,
    )
    _write_run(
        tmp_path,
        model=model,
        task="aftraj_audit",
        cases=100,
        score=0.2,
        input_tokens=1000,
        output_tokens=200,
        latency_s=50,
    )
    _write_run(
        tmp_path,
        model=model,
        task="aftraj_audit_test",
        cases=20,
        score=0.9,
        input_tokens=400,
        output_tokens=100,
        latency_s=20,
    )

    structure = [
        {
            "benchmark_source": "AgentCollabBench",
            "runnable_tasks": ["agent_collab_idr", "agent_collab_rtd"],
        },
        {
            "benchmark_source": "AFTraj-2K",
            "runnable_tasks": ["aftraj_audit", "aftraj_audit_test"],
        },
    ]
    runs = discover_task_runs(tmp_path)
    _, _, aggregates = aggregate_benchmarks(runs, benchmark_structure=structure)

    collab = aggregates[(model, "AgentCollabBench")]
    assert collab.score == 0.875
    assert collab.input_per_case == 25.0
    assert collab.output_per_case == 12.5
    assert collab.latency_per_case == 3.5

    aftraj = aggregates[(model, "AFTraj-2K")]
    assert aftraj.tasks == ("aftraj_audit",)
    assert aftraj.cases == 100
    assert aftraj.score == 0.2


def test_gaia_uses_levels_when_whole_validation_is_interrupted(tmp_path):
    model = "Qwen3-4B-Instruct-2507"
    _write_run(
        tmp_path,
        model=model,
        task="gaia_validation",
        cases=56,
        score=0.5,
        input_tokens=560,
        output_tokens=112,
        latency_s=56,
        expected=165,
        with_metrics=False,
    )
    for level, cases in ((1, 53), (2, 86), (3, 26)):
        _write_run(
            tmp_path,
            model=model,
            task=f"gaia_validation_level_{level}",
            cases=cases,
            score=level / 10,
            input_tokens=cases * 100,
            output_tokens=cases * 20,
            latency_s=cases * 2,
        )

    structure = [
        {
            "benchmark_source": "GAIA",
            "runnable_tasks": [
                "gaia_validation",
                "gaia_validation_level_1",
                "gaia_validation_level_2",
                "gaia_validation_level_3",
            ],
        }
    ]
    runs = discover_task_runs(tmp_path)
    _, _, aggregates = aggregate_benchmarks(runs, benchmark_structure=structure)
    gaia = aggregates[(model, "GAIA")]

    assert gaia.cases == 165
    assert gaia.tasks == (
        "gaia_validation_level_1",
        "gaia_validation_level_2",
        "gaia_validation_level_3",
    )
    assert gaia.status == "complete"


def test_workbook_has_grouped_benchmark_headers(tmp_path):
    model = "Qwen3-4B-Instruct-2507"
    _write_run(
        tmp_path,
        model=model,
        task="gsm8k",
        cases=10,
        score=0.8,
        input_tokens=1000,
        output_tokens=200,
        latency_s=30,
    )
    output = write_benchmark_workbook(tmp_path, tmp_path / "summary.xlsx")

    from openpyxl import load_workbook

    workbook = load_workbook(output, data_only=False)
    sheet = workbook["Benchmark Summary"]
    assert sheet["A1"].value == "Model"
    assert sheet["B1"].value == "GSM8K"
    assert [sheet.cell(2, column).value for column in range(2, 6)] == [
        "Acc.",
        "Input.",
        "Output.",
        "Lat.",
    ]
    assert sheet["A3"].value == model
    assert sheet["B3"].value == 0.8
